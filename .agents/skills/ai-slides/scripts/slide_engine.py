import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# =========================================================
# DESIGN SYSTEM PALETTE
# =========================================================
class PALETTE:
    BG_DARK = RGBColor(15, 23, 42)        # #0F172A (Deep Slate 900)
    CARD_BG = RGBColor(30, 41, 59)        # #1E293B (Card Slate 800)
    CARD_BG_ALT = RGBColor(39, 53, 73)    # #273549 (Elevated Slate)
    CARD_BORDER = RGBColor(51, 65, 85)    # #334155 (Subtle Border)
    TABLE_HEADER = RGBColor(21, 32, 48)   # Deep navy header

    PRIMARY = RGBColor(56, 189, 248)       # #38BDF8 (Sky Blue)
    ACCENT_EMERALD = RGBColor(52, 211, 153) # #34D399 (Emerald Green)
    ACCENT_AMBER = RGBColor(251, 191, 36)   # #FBBF24 (Gold / Amber)
    ACCENT_ROSE = RGBColor(244, 63, 94)     # #F43F5E (Rose / Crimson)

    TEXT_WHITE = RGBColor(248, 250, 252)   # #F8FAFC (Headings)
    TEXT_LIGHT = RGBColor(226, 232, 240)   # #E2E8F0 (Body)
    TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8 (Subtitles)

FONT_HEADING = "Segoe UI"
FONT_BODY = "Segoe UI"


class SlideDeck:
    """Reusable high-fidelity slide deck generator wrapping python-pptx."""

    def __init__(self, width_in=13.333333, height_in=7.5):
        self.prs = Presentation()
        self.prs.slide_width = Inches(width_in)
        self.prs.slide_height = Inches(height_in)
        self.width_in = width_in
        self.height_in = height_in

    def add_slide(self, badge="", title="", subtitle="", badge_color=PALETTE.PRIMARY):
        """Creates a dark-theme slide with a standardized modern header."""
        blank_layout = self.prs.slide_layouts[6]
        slide = self.prs.slides.add_slide(blank_layout)

        # Full-bleed background
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, Inches(self.width_in), Inches(self.height_in)
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = PALETTE.BG_DARK
        bg.line.fill.background()

        # Badge
        if badge:
            bbox = slide.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(11.733), Inches(0.28))
            btf = bbox.text_frame
            btf.word_wrap = True
            btf.margin_left = btf.margin_top = btf.margin_right = btf.margin_bottom = 0
            bp = btf.paragraphs[0]
            bp.text = badge.upper()
            bp.font.name = FONT_HEADING
            bp.font.size = Pt(9.5)
            bp.font.bold = True
            bp.font.color.rgb = badge_color

        # Title
        if title:
            tbox = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.733), Inches(0.55))
            ttf = tbox.text_frame
            ttf.word_wrap = True
            ttf.margin_left = ttf.margin_top = ttf.margin_right = ttf.margin_bottom = 0
            tp = ttf.paragraphs[0]
            tp.text = title
            tp.font.name = FONT_HEADING
            tp.font.size = Pt(21)
            tp.font.bold = True
            tp.font.color.rgb = PALETTE.TEXT_WHITE

        # Subtitle
        if subtitle:
            sbox = slide.shapes.add_textbox(Inches(0.8), Inches(1.24), Inches(11.733), Inches(0.35))
            stf = sbox.text_frame
            stf.word_wrap = True
            stf.margin_left = stf.margin_top = stf.margin_right = stf.margin_bottom = 0
            sp = stf.paragraphs[0]
            sp.text = subtitle
            sp.font.name = FONT_BODY
            sp.font.size = Pt(11.5)
            sp.font.color.rgb = PALETTE.TEXT_MUTED

        # Divider line
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.65), Inches(11.733), Inches(0.015)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = PALETTE.CARD_BORDER
        line.line.fill.background()

        return slide

    def add_card(self, slide, left, top, width, height, bg_color=PALETTE.CARD_BG, border_color=PALETTE.CARD_BORDER, border_width_pt=1):
        """Adds a rounded rectangle card container."""
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(border_width_pt)
        else:
            card.line.fill.background()
        return card

    def add_table(self, slide, rows, cols, left, top, width, height):
        """Adds an unstyled table ready for custom formatting."""
        table_shape = slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(height))
        return table_shape.table

    def format_cell(self, cell, text, font_size=10, bold=False, color=PALETTE.TEXT_LIGHT, bg_color=None, align=PP_ALIGN.LEFT):
        """Formats a table cell with padding and typography."""
        if bg_color:
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_color
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.12)
        tf.margin_right = Inches(0.12)
        tf.margin_top = Inches(0.06)
        tf.margin_bottom = Inches(0.06)
        p = tf.paragraphs[0]
        p.text = str(text)
        p.alignment = align
        p.font.name = FONT_BODY
        p.font.size = Pt(font_size)
        p.font.bold = bold
        p.font.color.rgb = color

    def save(self, output_path):
        """Saves the presentation."""
        self.prs.save(output_path)
        print(f"Presentation saved successfully to: {os.path.abspath(output_path)}")
